"""Load the sets and parameters used by the transportation model."""

from __future__ import annotations

from dataclasses import dataclass  # Creates a class for storing related data.
from pathlib import Path          # Represents the Excel file's path.

import numpy as np                      # Stores numerical vectors and matrices.
from openpyxl import load_workbook       # Reads the Excel workbook.


@dataclass
class Problem:
    """Defining the sets and parameters of the transportation model."""

    # Defining sets of the model
    centres: tuple[str, ...]     # L: logistics centres, having m = number of centres
    regions: tuple[str, ...]     # R: affected regions, having n = number of regions
    resources: tuple[str, ...]   # V: resource types, having p = number of resources

    # Defining parameters of the model
    centre_xy: np.ndarray    # Shape (m, 2): each centre's [x, y] coordinates.
    region_xy: np.ndarray    # Shape (n, 2): each region's [x, y] coordinates.
    supply: np.ndarray       # S[i, k], shape (m, p): available stock.
    demand: np.ndarray       # D[j, k], shape (n, p): minimum requirements.
    capacity: np.ndarray     # C[i], length m: equivalent handling capacity.
    handling: np.ndarray     # w[k], length p: handling effort per resource unit.
    rate: np.ndarray         # r[k], length p: transport cost in euros/unit/km.
    trucks_available: np.ndarray  # Descriptive information only; not used by the LP.
    truck_capacity: np.ndarray    # Descriptive information only; not used by the LP.
    cargo_volume: np.ndarray      # v[k]: space occupied by one resource unit.

    fleet_capacity: np.ndarray    # F[i]: total cargo-space limit read directly from Excel.


def read_data(path: Path) -> Problem:
    """Read the Excel tables and return a populated Problem object."""
    workbook = load_workbook(path, read_only=True, data_only=True)

    def table(sheet):
        # Read cell values. The first row supplies the column names.
        rows = workbook[sheet].iter_rows(values_only=True)
        header = next(rows)

        # Pair each column name with its value to create one dictionary per row.
        # For example: {"centre_id": "L1", "handling_capacity": 690, ...}.
        records = []
        for row in rows:
            if any(value is not None for value in row):  # Skip empty rows.
                records.append(dict(zip(header, row)))
        return records

    try:
        # Each variable below holds a list of dictionaries from one Excel sheet.
        centres = table("Centres")
        regions = table("Regions")
        resources = table("Resources")
        supply = table("Supply")
        demand = table("Demand")
    finally:
        # The values above stay in memory after the workbook is closed.
        workbook.close()

    # Read IDs exactly as provided in the workbook, including V1-V5.
    # Their order determines the positions in every array.
    # Python positions start at 0: centre 0 is L1, resource 0 is V1.
    centre_ids = tuple(row["centre_id"] for row in centres)
    region_ids = tuple(row["region_id"] for row in regions)
    resource_ids = tuple(row["resource_id"] for row in resources)

    def matrix(rows, origin_column, origins, value_column):
        # Convert the long Excel table into a lookup keyed by (location, resource).
        # For supply, an example entry is ("L1", "V1"): 230.
        values = {(row[origin_column], row["resource_id"]): row[value_column] for row in rows}

        # Build a matrix with one row per location and one column per resource.
        # float allows fractional quantities, as required by the continuous LP.
        return np.array([[values[origin, resource_id] for resource_id in resource_ids]
                         for origin in origins], dtype=float)

    # Fill the Problem fields. Named arguments show which input each array holds.
    problem = Problem(
        centres=centre_ids,
        regions=region_ids,
        resources=resource_ids,
        centre_xy=np.array([[row["x_km"], row["y_km"]] for row in centres], dtype=float),
        region_xy=np.array([[row["x_km"], row["y_km"]] for row in regions], dtype=float),
        supply=matrix(supply, "centre_id", centre_ids, "available_units"),
        demand=matrix(demand, "region_id", region_ids, "minimum_units"),
        capacity=np.array([row["handling_capacity"] for row in centres], dtype=float),
        handling=np.array([row["handling_per_unit"] for row in resources], dtype=float),
        rate=np.array([row["transport_rate_eur_per_unit_km"] for row in resources], dtype=float),
        trucks_available=np.array([row["trucks_available"] for row in centres], dtype=float),
        truck_capacity=np.array([row["truck_capacity"] for row in centres], dtype=float),
        cargo_volume=np.array([row["cargo_volume_per_unit"] for row in resources], dtype=float),
        fleet_capacity=np.array([row["fleet_capacity"] for row in centres], dtype=float),
    )
    # Fleet capacity and cargo volume must use the same space unit.
    # Truck counts and per-truck capacities are descriptive, so no product check is needed.
    for values in (problem.fleet_capacity, problem.cargo_volume):
        if not np.isfinite(values).all() or (values < 0).any():
            raise ValueError("Fleet capacity and cargo volumes must be finite and nonnegative")
    # main.py receives this object and passes it to the solver and analysis.
    return problem
